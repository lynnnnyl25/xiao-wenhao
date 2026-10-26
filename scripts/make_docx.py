# -*- coding: utf-8 -*-
"""生成《小问号》教育智能体比赛申报书"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# ---------- 全局字体 ----------
st = doc.styles['Normal']
st.font.name = '微软雅黑'
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')


def set_font(run, name='微软雅黑', size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def h(text, level=1):
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    set_font(r, size=16 if level == 0 else 13, bold=True, color=(0x2E, 0x2A, 0x27))
    return p


def para(text, size=10.5, bold=False, align=None, color=None, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p


def bullet(text, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=size)
    return p


def table(rows, widths=None, header=True):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cell = cells[ci]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            set_font(r, size=9.5, bold=(header and ri == 0))
            if widths:
                cell.width = Cm(widths[ci])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


# ============================================================ 封面
para('教育智能体比赛 · 作品申报书', size=11, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, color=(0x8A, 0x81, 0x7C))
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('小问号')
set_font(r, size=32, bold=True, color=(0xFF, 0x8A, 0x5B))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('苏格拉底 × 费曼 主动学习智能体')
set_font(r, size=15, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('—— 让 AI 从「替学生回答」转向「引导学生学会思考」')
set_font(r, size=11, color=(0x8A, 0x81, 0x7C))
doc.add_paragraph()

table([
    ['作品名称', '小问号 · 苏格拉底 × 费曼 主动学习智能体'],
    ['应用学段/学科', '小学 3-6 年级 · 信息科技（对标《义务教育信息科技课程标准（2022年版）》）'],
    ['技术底座', 'DeepSeek（deepseek-chat / deepseek-reasoner）'],
    ['交付形态', 'Web 应用（公网可访问）+ 开源仓库 + 可复用教学内核 Skill'],
    ['核心方法', '苏格拉底式追问 + 费曼学习法 + 认知缺口诊断'],
], widths=[4, 12], header=False)

doc.add_page_break()

# ============================================================ 一、问题
h('一、要解决的教学问题', 1)

para('1.1 假性掌握：AI 越会答，学生越学不到', bold=True)
para(
    '学生问 AI「什么是变量」，AI 立刻给出一段结构完整、措辞精准的解释。学生读完点头，'
    '觉得自己懂了——但换一道题还是不会。当前主流 AI 助教的产品逻辑是「问→答」，'
    '它在最大化"答案的易得性"，而这恰恰在最小化"理解的深度"。'
    '认知科学上，这是一个已被反复验证的现象：流畅阅读产生的熟悉感会被误认为掌握（illusion of knowing）。'
)
para(
    '本作品要解决的正是这一点：把 AI 从"答案的提供者"改造成"思考的引导者"。',
    bold=True,
)

para('1.2 看不见的缺口：教师不知道学生卡在哪', bold=True)
para(
    '学生在哪里没懂，通常是模糊的印象。一次对话结束后，教师拿到的只有聊天记录，'
    '无法回答三个关键问题：这个孩子到底卡在哪个环节？他是背下来了还是真懂了？下一课该从哪里补起？'
)

para('1.3 小学场景的特殊困难', bold=True)
bullet('8-11 岁儿童打字慢、表达能力有限，无法用长段文字描述自己的困惑；')
bullet('抽象术语对他们不生效，"变量""作用域""递归"等词本身就构成理解障碍；')
bullet('连续追问极易造成挫败，孩子会直接放弃。')

# ============================================================ 二、方案
h('二、解决方案：G-Loop 引导闭环', 1)

para(
    '小问号的核心是一个名为 G-Loop（Gap Loop）的教学状态机。'
    '它不管理"对话"，而是管理"缺口"——每一轮的目标不是把话接下去，而是把下一个缺口逼出来。'
)
para(
    '提问（苏格拉底）→ 学生回答 → 诊断缺口 → 更新图谱 → 决定下一问 ↑ ‖ '
    '讲不清 → 回到追问 ← 费曼环节（学生讲给 AI 听）← 疑似懂了',
    size=10, color=(0x8A, 0x81, 0x7C), align=WD_ALIGN_PARAGRAPH.CENTER,
)

para('2.1 苏格拉底追问引擎', bold=True)
para('按 6 层阶梯推进，不跳级；学生在哪一层卡住就在那一层反复，通了才往上登。')
table([
    ['层级', '动作', '话术示例'],
    ['L1', '唤起经验', '「你平时……的时候，会怎么样？」'],
    ['L2', '概念澄清', '「你说的 XX，具体是指什么？」'],
    ['L3', '举例反例', '「有没有一个时候，它就不灵了？」'],
    ['L4', '类比迁移', '「它像你生活里的什么东西？为什么像？」'],
    ['L5', '追问原因', '「你为什么会这么想？你是怎么知道的？」'],
    ['L6', '归纳表达', '「用一句话说说，它到底是什么？」'],
], widths=[1.8, 3.2, 11])

para('2.2 费曼环节：让学生教 AI', bold=True)
para(
    '当学生疑似掌握，AI 切换到「从没学过的同桌」角色，请学生讲给它听。'
    'AI 只在含糊、术语依赖、逻辑断链处打断，对应费曼技巧第三步「回到材料」。'
    '通过需三条全满足：无未定义术语 + 有生活类比 + 能回答变式追问。'
)

para('2.3 认知缺口图谱', bold=True)
para('把学生的每一句话结构化成缺口事件，归入小学版 7 类，每条都带学生原话引用。')
table([
    ['类型', '含义', '小学生典型表现（真实样例）'],
    ['定义缺失', '说不出"是什么"', '「就是……电脑里面放东西的地方吧」'],
    ['机制不清', '说不出"怎么工作"', '「因为……加了就更难猜？」'],
    ['类比失效', '打比方打错地方', '把"放进去"当成"加上去"，算出 5+3=8'],
    ['术语依赖', '用没懂的词解释词', '「电脑有内存，就记住了」'],
    ['因果想当然', '把先后当因果', '「因为电脑是先发明 0 和 1 的」'],
    ['边界模糊', '不知何时不适用', '「密码加生日更安全，因为更长了」'],
    ['迁移失败', '换个情境就不会', '会做这道，换个说法就不会'],
], widths=[2.4, 4.2, 9.4])

para(
    '每条缺口记录发现轮次、闭合轮次与学生原话，形成可追溯的证据链，'
    '并实时渲染为雷达图、理解度成长曲线与逐条证据链。'
)

# ============================================================ 三、创新点
h('三、核心创新点', 1)

para('创新点一：把"没懂"变成可追踪、可量化、可闭合的结构化资产', bold=True)
para(
    '现有苏格拉底式 AI 导师多为"一次一提示"的单引擎实现：给出提示，等待回答，再给下一个提示。'
    '学生的困难在其中是瞬时的、不落地的。'
    '小问号把每一次"说不清"沉淀为带证据的缺口事件，并在后续轮次中主动闭合它，'
    '使"学习进展"第一次可以被度量（闭合率），而不是被感觉。'
)

para('创新点二：拒答护栏——把"AI 不给答案"变成可当场验证的数字', bold=True)
para(
    '"不给答案"很容易承诺，很难证明。本作品内置护栏，在对话过程中实时统计三项指标，'
    '评委可在使用时随时查看，无需信任开发者的自述：'
)
table([
    ['指标', '定义', '目标', '本项目实测'],
    ['答问比 Q/A', '疑问句数 ÷ 结论陈述句数', '≥ 3', '60 – 100'],
    ['泄露率 Leak Rate', '直接给出答案的轮次占比', '≤ 5%', '0%'],
    ['缺口闭合率', '已闭合缺口 ÷ 诊断出的缺口总数', '≥ 60%', '100%'],
], widths=[3.2, 6.4, 2.6, 3.8])
para(
    '数据来源：pnpm selfcheck，在 5 段回放脚本共 39 轮对话上自动统计，结果可复现。'
    '护栏通过率 39/39 = 100%。',
    size=9.5, color=(0x8A, 0x81, 0x7C),
)

para('创新点三：面向小学的儿童认知适配', bold=True)
bullet('儿童语言公约：一句话 ≤ 25 字，术语必须翻译成动作（"变量"→"贴了名字的盒子"）；')
bullet('三个状态按钮（我不知道 / 给点提示 / 我讲完了），绕开打字门槛；')
bullet('挫败防护：连续答不上来自动降层重问，只给"最小提示"，绝不给完整答案；')
bullet('内容对标《义务教育信息科技课程标准（2022 年版）》，覆盖算法与程序、数据与编码、网络与安全、人工智能初步四大板块。')

# ============================================================ 四、技术
h('四、技术实现', 1)

para('4.1 架构', bold=True)
para(
    'Vite + React 18 + TypeScript + Tailwind 静态前端，浏览器直连 DeepSeek API。'
    'API Key 仅存于用户本地 localStorage，不经任何中间服务器——'
    '对小学生的对话数据而言，这是隐私合规上更稳妥的选择，同时免去后端运维与部署成本。'
)

para('4.2 三角色协同（同一模型，靠 Prompt 与结构化输出切换）', bold=True)
table([
    ['角色', '职责', '结构化输出'],
    ['Socratic Engine', '生成单个追问，按阶梯升降层', '{say, level, target_gap, intent, give_up_hint}'],
    ['Gap Diagnoser', '从学生原话识别 7 类缺口', '{gaps[{type,evidence,confidence}], understood_level}'],
    ['Feynman Reviewer', '扮演不懂的同桌，只在断链处打断', '{say, passed, stuck_at, missing, praise}'],
], widths=[3.6, 6.4, 6])

para('4.3 体验模式（评委审核兜底）', bold=True)
para(
    '内置 5 段完整回放脚本（变量 / 密码安全 / 二进制 / 循环 / 人工智能初步），'
    '不需要 API Key、不依赖网络模型即可完整体验 G-Loop 全流程与图谱动态生长，'
    '确保评委在任何网络与账号条件下都能完成审核。'
)

para('4.4 与已有开源项目的差异', bold=True)
para(
    '开源界已有苏格拉底提问与费曼技巧的单引擎实现'
    '（socratic-workshop、m4vic/socratic、feynman-tutor 等）。本作品的差异在于：'
)
bullet('缺口结构化：把"没懂"落成 7 类可追踪、可量化、可闭合的缺口事件；')
bullet('护栏可审计：用三项指标自证"真的不给答案"，评委可当场验证；')
bullet('场景落地：面向小学信息科技，含儿童语言公约与挫败防护，对标 2022 义教新课标。')

# ============================================================ 五、应用
h('五、应用价值与推广', 1)
table([
    ['角色', '价值'],
    ['学生', '被带着想懂而非被告诉答案；学习护照记录"我今天想通了什么"，形成成就感'],
    ['教师', '看到班级高频缺口 Top5 与典型迷思原话，直接决定下一课从哪里切入'],
    ['家长', '区分"背下来了"与"真懂了"，看懂孩子的思维过程而非仅看对错'],
    ['教研', '7 类缺口分类法可作为小学生概念理解困难的观察框架'],
], widths=[2.4, 13.6])

para('后续规划：语音输入与朗读、班级端多学生看板、间隔复习（SM-2）自动排程、校本题库接入。')

# ============================================================ 六、材料
doc.add_page_break()
h('六、评审材料与链接', 1)
para('（部署完成后填入下表；体验模式无需 API Key，评委可直接点开使用）')
table([
    ['材料', '说明', '链接'],
    ['在线体验（主链接）', '公网可访问，默认进入体验模式，免 API Key', 'https://lynnnnyl25.github.io/xiao-wenhao/'],
    ['GitHub 仓库', '源码、教学内核 Prompt、评测脚本、部署说明', 'https://github.com/lynnnnyl25/xiao-wenhao'],
    ['演示视频', '3 分钟演示，含小学生真实使用片段', '待补充'],
    ['技术文档', '仓库 README 与 docs/ 目录', 'https://github.com/lynnnnyl25/xiao-wenhao#readme'],
], widths=[3.4, 8.4, 4.2])

para('推荐审核路径（约 5 分钟）：', bold=True)
bullet('打开在线体验链接，默认进入体验模式；')
bullet('点「播放完整过程」，观察 AI 全程只问不答、右侧缺口图谱随之生长；')
bullet('查看右侧「拒答护栏自检」面板的三项指标；')
bullet('点「生成学习护照」查看本次学习的结构化产出；')
bullet('切到「真模型」填入自己的 DeepSeek Key，可与真实模型自由对话。')

doc.add_paragraph()
para('技术栈：Vite · React 18 · TypeScript · Tailwind · DeepSeek API',
     size=9, color=(0x8A, 0x81, 0x7C), align=WD_ALIGN_PARAGRAPH.CENTER)
para('许可证：MIT', size=9, color=(0x8A, 0x81, 0x7C), align=WD_ALIGN_PARAGRAPH.CENTER)

out = '小问号-教育智能体申报书.docx'
doc.save(out)
print('saved:', out)
