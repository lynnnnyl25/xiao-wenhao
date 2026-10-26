# -*- coding: utf-8 -*-
"""校验申报书内容"""
from docx import Document
import glob

f = glob.glob("*.docx")[0]
d = Document(f)
print("file:", f)
print("paragraphs:", len(d.paragraphs), "tables:", len(d.tables))

n = 0
for p in d.paragraphs:
    t = p.text.strip()
    if t:
        n += 1
        if n <= 45:
            print(f"[{n}] {t[:95]}")
print("non-empty paragraphs:", n)

for i, t in enumerate(d.tables):
    print(f"--- table {i}: {len(t.rows)}x{len(t.columns)} | first row: {[c.text for c in t.rows[0].cells]}")
